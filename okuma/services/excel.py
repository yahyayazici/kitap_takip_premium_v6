"""Raporların Excel çıktısı için ortak yardımcı."""

import io

import openpyxl
from django.http import HttpResponse
from openpyxl.styles import Alignment, Font, PatternFill

BASLIK_DOLGU = PatternFill("solid", fgColor="1D1D1F")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")


def excel_yaniti(dosya_adi: str, sayfalar: list[tuple[str, list[str], list[list]]]) -> HttpResponse:
    """sayfalar: [(sayfa adı, başlıklar, satırlar), ...]"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for ad, basliklar, satirlar in sayfalar:
        ws = wb.create_sheet(title=ad[:31] or "Sayfa")
        ws.append(basliklar)
        for hucre in ws[1]:
            hucre.fill = BASLIK_DOLGU
            hucre.font = BASLIK_YAZI
            hucre.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for satir in satirlar:
            ws.append([round(v, 2) if isinstance(v, float) else v for v in satir])
        for kolon in ws.columns:
            uzunluk = max(len(str(c.value or "")) for c in kolon)
            ws.column_dimensions[kolon[0].column_letter].width = max(7, min(48, uzunluk + 2))
        ws.freeze_panes = "B2"
    cikti = io.BytesIO()
    wb.save(cikti)
    yanit = HttpResponse(cikti.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    yanit["Content-Disposition"] = f'attachment; filename="{dosya_adi}.xlsx"'
    return yanit
