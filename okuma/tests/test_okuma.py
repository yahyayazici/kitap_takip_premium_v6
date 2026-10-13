import datetime
import io
from decimal import Decimal
from pathlib import Path

import openpyxl
from django.conf import settings
from django.contrib.auth.models import User
from django.test import TestCase, override_settings

from okuma.models import Deneme, OptikSablon, OptikSatir, Soru
from okuma.services import analiz
from okuma.services.anahtar_excel import AnahtarHatasi, anahtar_yukle
from okuma.services.lgs import LGS_KATSAYILARI
from okuma.services.optik import optik_yukle, satir_ayristir
from okuma.services.puanlama import deneme_hesapla, kitapcik_tahmin_et, net_hesapla

ORNEK = Path(settings.BASE_DIR) / "ornek_veri"

# Testler optik alan adının URL yapılandırmasıyla, kullanıcı yönetimi açık ve
# yalnız personel erişimiyle çalışır (ana sitenin rol sistemi burada devreye girmez).
OPTIK = override_settings(ROOT_URLCONF="config.optik_urls", OKUMA_ERISIM=None, OKUMA_KULLANICI_YONETIMI=True)

# İki test: Türkçe 4 soru, Matematik 3 soru. B kitapçığında sıralar ters.
ANAHTAR = [
    ("Sözel (TÜR)", "Türkçe", 1, 4, "A", "T.1", "Sözcükte anlam"),
    ("Sözel (TÜR)", "Türkçe", 2, 3, "B", "T.2", "Paragraf"),
    ("Sözel (TÜR)", "Türkçe", 3, 2, "C", "T.2", "Paragraf"),
    ("Sözel (TÜR)", "Türkçe", 4, 1, "D", "", "Yazım"),
    ("Sayısal (MAT)", "Matematik", 1, 3, "B", "M.1", "Rasyonel sayılar"),
    ("Sayısal (MAT)", "Matematik", 2, 2, "C", "M.1", "Rasyonel sayılar"),
    ("Sayısal (MAT)", "Matematik", 3, 1, "A", "M.2", "Oran orantı"),
]
# A sırasıyla anahtar: ABCD BCA  → B kitapçığında: DCBA ACB


def excel_uret(satirlar=ANAHTAR) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ÇS", None, "7.SINIF DENEME"])
    ws.append([])
    ws.append(["Kitapçık", "Test", "Ders", "A Soru", "B Soru", "Cevap", "Kazanım Kodu", "Kazanım-1", "Kazanım-2"])
    for test, ders, a, b, cevap, kod, kaz in satirlar:
        ws.append(["A", test, ders, a, b, cevap, kod, kaz, None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def satir_uret(numara: str, kitapcik: str, cevaplar: str, sinif="7", sube="A", ad="") -> str:
    # Çinili Saray şablonu: numara 20–25, kitapçık 26, sınıf 27, şube 28, ad 29–50, cevap 51+
    return "0" + " " * 18 + numara.rjust(6) + kitapcik + sinif + sube + ad.ljust(22) + cevaplar


@OPTIK
class TemelTest(TestCase):
    def setUp(self):
        self.sablon = OptikSablon.objects.get(ad__startswith="Çinili Saray")
        self.deneme = Deneme.objects.create(
            ad="Test deneme", tarih=datetime.date(2026, 10, 10), sinif_seviyesi="7", sablon=self.sablon
        )
        anahtar_yukle(self.deneme, excel_uret())

    def yukle(self, *satirlar, **kw):
        veri = "\r\n".join(satirlar).encode("cp1254")
        return optik_yukle(self.deneme, self.sablon, "t.dat", veri, **kw)


@OPTIK
class SatirAyristirmaTest(TemelTest):
    def test_alanlar(self):
        o = satir_ayristir(1, satir_uret("00048", "A", "ABCDBC*", ad="AHMET ŞAHİN"), self.sablon, 7)
        self.assertEqual((o.numara, o.kitapcik, o.sinif, o.sube, o.ad_soyad), ("48", "A", "7", "A", "AHMET ŞAHİN"))
        self.assertEqual(o.cevaplar, "ABCDBC*")
        self.assertFalse(o.hatalar)

    def test_numara_alaninda_fazladan_isaret(self):
        satir = "0" + " " * 18 + "5   57" + "B  " + " " * 22 + "ABCDBCA"
        o = satir_ayristir(1, satir, self.sablon, 7)
        self.assertEqual(o.numara, "57")
        self.assertTrue(o.uyarilar)

    def test_kisa_satir_bos_ile_doldurulur(self):
        o = satir_ayristir(1, satir_uret("12", "A", "AB"), self.sablon, 7)
        self.assertEqual(o.cevaplar, "AB     ")


@OPTIK
class AnahtarTest(TemelTest):
    def test_yapi(self):
        testler = list(self.deneme.testler.all())
        self.assertEqual([(t.ders, t.oturum, t.soru_sayisi, t.lgs_ders, int(t.katsayi)) for t in testler], [
            ("Türkçe", "Sözel", 4, "turkce", 4),
            ("Matematik", "Sayısal", 3, "matematik", 4),
        ])
        self.assertEqual(self.deneme.kitapcik_listesi, ["A", "B"])
        self.assertEqual(Soru.objects.get(test__ders="Türkçe", no=1).kitapcik_siralari.get().no, 4)

    def test_eksik_b_sirasi_reddedilir(self):
        bozuk = list(ANAHTAR)
        bozuk[0] = ("Sözel (TÜR)", "Türkçe", 1, 3, "A", "", "x")  # B=3 iki kez
        with self.assertRaises(AnahtarHatasi):
            anahtar_yukle(self.deneme, excel_uret(bozuk))


@OPTIK
class PuanlamaTest(TemelTest):
    def test_net_kurali(self):
        self.assertEqual(net_hesapla(10, 3, 3), Decimal("9.00"))
        self.assertEqual(net_hesapla(10, 4, 4), Decimal("9.00"))
        self.assertEqual(net_hesapla(0, 6, 3), Decimal("-2.00"))  # negatif net kırpılmaz
        self.assertEqual(net_hesapla(5, 5, 0), Decimal("5.00"))

    def test_b_kitapcigi_a_sirasina_cevrilir(self):
        self.yukle(satir_uret("1", "A", "ABCDBCA"), satir_uret("2", "B", "DCBAACB"))
        deneme_hesapla(self.deneme)
        for numara in ("1", "2"):
            s = self.deneme.sonuclar.get(ogrenci__numara=numara)
            self.assertEqual((s.dogru, s.yanlis, s.bos), (7, 0, 0))
            self.assertEqual(s.cevaplar_a, "ABCDBCA")

    def test_yanlis_bos_cift_isaret(self):
        # Türkçe: A doğru, B→C yanlış, boş, * (çift) yanlış → 1D 2Y 1B → net 1-2/3
        self.yukle(satir_uret("1", "A", "AC *BCA"))
        deneme_hesapla(self.deneme)
        tr = self.deneme.sonuclar.get().test_sonuclari.get(test__ders="Türkçe")
        self.assertEqual((tr.dogru, tr.yanlis, tr.bos, tr.net), (1, 2, 1, Decimal("0.33")))

    def test_iptal_soru_degerlendirme_disi(self):
        self.yukle(satir_uret("1", "A", "DBCDBCA"))  # 1. soru yanlış
        soru = Soru.objects.get(test__ders="Türkçe", no=1)
        soru.iptal = True
        soru.save()
        deneme_hesapla(self.deneme)
        s = self.deneme.sonuclar.get()
        self.assertEqual((s.dogru, s.yanlis, s.bos), (6, 0, 0))

    def test_ek_dogru_sik(self):
        soru = Soru.objects.get(test__ders="Türkçe", no=1)
        soru.ek_dogrular = "B"
        soru.save()
        self.yukle(satir_uret("1", "A", "BBCDBCA"))
        deneme_hesapla(self.deneme)
        self.assertEqual(self.deneme.sonuclar.get().dogru, 7)

    def test_lgs_tablo_puani(self):
        self.yukle(satir_uret("1", "A", "ABCDBCA"))
        deneme_hesapla(self.deneme)
        t = LGS_KATSAYILARI[2022]
        # Tam net, LGS soru sayısına oranlanır: Türkçe 4/4 → 20/20, Matematik 3/3 → 20/20.
        beklenen = Decimal(str(t["taban"])) + 20 * Decimal(str(t["turkce"])) + 20 * Decimal(str(t["matematik"]))
        self.assertEqual(self.deneme.sonuclar.get().puan, beklenen.quantize(Decimal("0.001")))

    def test_lgs_2024_ornek_hesap(self):
        # Araştırma notlarındaki doğrulama örneği: 15/8/9/7/12/14 net → ≈409,17
        t = LGS_KATSAYILARI[2024]
        puan = t["taban"] + 15 * t["turkce"] + 8 * t["inkilap"] + 9 * t["din"] + 7 * t["yabanci_dil"] + 12 * t["matematik"] + 14 * t["fen"]
        self.assertAlmostEqual(puan, 409.166, places=3)
        for yil, k in LGS_KATSAYILARI.items():
            tam = k["taban"] + 20 * (k["turkce"] + k["matematik"] + k["fen"]) + 10 * (k["inkilap"] + k["din"] + k["yabanci_dil"])
            self.assertAlmostEqual(tam, 500, delta=0.5, msg=yil)  # katsayılar 3 haneye yuvarlı

    def test_2025_katsayilari_kurum_ornegi(self):
        # Kurumsal Deneme-1 birincisi: Tr 19D1Y, İnk 10D, Din 7D3Y, İng 9D1Y, Mat 18D2Y, Fen 18D2Y → ≈461
        t = LGS_KATSAYILARI[2025]
        netler = {"turkce": 19 - 1 / 3, "inkilap": 10, "din": 7 - 1, "yabanci_dil": 9 - 1 / 3, "matematik": 18 - 2 / 3, "fen": 18 - 2 / 3}
        puan = t["taban"] + sum(t[k] * v for k, v in netler.items())
        self.assertAlmostEqual(puan, 460.72, places=1)

    def test_2022_kurum_okuma_sitesiyle_ayni(self):
        # XYS sıralı listesi: Ahmet Sevban Taşdan 463,100 — Hüseyin Kemal İnce 457,180
        t = LGS_KATSAYILARI[2022]
        for netler, beklenen in [
            ({"turkce": 19 - 1 / 3, "inkilap": 10, "din": 6, "yabanci_dil": 9 - 1 / 3, "matematik": 18 - 2 / 3, "fen": 18 - 2 / 3}, 463.100),
            ({"turkce": 18 - 2 / 3, "inkilap": 10, "din": 10, "yabanci_dil": 6, "matematik": 19 - 1 / 3, "fen": 16 - 4 / 3}, 457.180),
        ]:
            self.assertAlmostEqual(t["taban"] + sum(t[k] * v for k, v in netler.items()), beklenen, delta=0.2)

    def test_deneme_ici_puan_100_500(self):
        self.deneme.puan_modu = Deneme.PuanModu.DENEME_ICI
        self.deneme.save()
        self.yukle(satir_uret("1", "A", "ABCDBCA"), satir_uret("2", "A", "ABC DDD"), satir_uret("3", "A", "       "))
        deneme_hesapla(self.deneme)
        puanlar = dict(self.deneme.sonuclar.values_list("ogrenci__numara", "puan"))
        self.assertEqual(puanlar["1"], Decimal("500.000"))
        self.assertEqual(puanlar["3"], Decimal("100.000"))

    def test_esit_puan_ayni_sira(self):
        self.deneme.puan_modu = Deneme.PuanModu.YOK
        self.deneme.save()
        self.yukle(
            satir_uret("1", "A", "ABCDBCA"), satir_uret("2", "B", "DCBAACB"), satir_uret("3", "A", "ABCD   ", sube="B")
        )
        deneme_hesapla(self.deneme)
        siralar = dict(self.deneme.sonuclar.values_list("ogrenci__numara", "sira_kurum"))
        self.assertEqual(siralar, {"1": 1, "2": 1, "3": 3})
        self.assertEqual(self.deneme.sonuclar.get(ogrenci__numara="3").sira_sube, 1)


@OPTIK
class EslestirmeTest(TemelTest):
    def test_kitapcik_bos_ise_tahmin(self):
        cevap = "DCBAACB" * 2  # tahmin için en az 10 doğru gerekir; test küçük olduğundan doğrudan çağır
        self.assertIsNone(kitapcik_tahmin_et(self.deneme, cevap[:7]))  # 7 doğru < eşik

    def test_tekrarlanan_numara_hesaba_katilmaz(self):
        y = self.yukle(satir_uret("1", "A", "ABCDBCA"), satir_uret("1", "A", "DDDDDDD"))
        ikinci = y.satirlar.get(satir_no=2)
        self.assertEqual(ikinci.durum, OptikSatir.Durum.HATA)
        self.assertFalse(ikinci.dahil)
        deneme_hesapla(self.deneme)
        self.assertEqual(self.deneme.sonuclar.get().dogru, 7)

    def test_otomatik_ogrenci_kapali(self):
        y = self.yukle(satir_uret("99", "A", "ABCDBCA"), otomatik_ogrenci=False)
        self.assertEqual(y.satirlar.get().durum, OptikSatir.Durum.HATA)


@OPTIK
class AnalizTest(TemelTest):
    def test_kazanim_orani(self):
        self.yukle(satir_uret("1", "A", "ABCDBCA"), satir_uret("2", "A", "AXXDDDD"))
        deneme_hesapla(self.deneme)
        rapor = {k.kazanim.kod or k.kazanim.metin: k for k in analiz.kazanim_raporu(self.deneme)}
        self.assertEqual(rapor["T.2"].oran, 50.0)  # 2 soru × 2 öğrenci, 2 doğru
        self.assertEqual(rapor["M.1"].oran, 50.0)
        self.assertEqual(rapor["T.1"].oran, 100.0)

    def test_madde_analizi(self):
        self.yukle(*[satir_uret(str(i), "A", "ABCDBCA" if i < 3 else "DDDDDDD") for i in range(1, 7)])
        deneme_hesapla(self.deneme)
        maddeler, kr20 = analiz.madde_analizi(self.deneme)
        ilk = maddeler[0]
        self.assertEqual(ilk.p, 0.33)
        self.assertEqual(ilk.r, 1.0)
        self.assertIn("Türkçe", kr20)


@override_settings(OKUMA_API_TOKEN="gizli")
@OPTIK
class ArayuzTest(TemelTest):
    def setUp(self):
        super().setUp()
        self.client.force_login(User.objects.create_user("y", is_staff=True))
        self.yukle(satir_uret("1", "A", "ABCDBCA"))
        deneme_hesapla(self.deneme)

    def test_sayfalar(self):
        pk = self.deneme.pk
        s = self.deneme.sonuclar.get()
        for url in ["/", f"/deneme/{pk}/", f"/deneme/{pk}/satirlar/", f"/deneme/{pk}/sonuclar/",
                    f"/deneme/{pk}/sonuclar.xlsx", f"/deneme/{pk}/karne/{s.pk}/", f"/deneme/{pk}/kazanim/",
                    f"/deneme/{pk}/madde/"]:
            self.assertEqual(self.client.get(url).status_code, 200, url)

    def test_giris_gerekli(self):
        self.client.logout()
        self.assertContains(self.client.get("/"), "Giriş yapın")
        self.assertEqual(self.client.get(f"/deneme/{self.deneme.pk}/").status_code, 302)

    def test_api_token(self):
        self.client.logout()
        url = f"/api/v1/denemeler/{self.deneme.pk}/sonuclar/"
        self.assertEqual(self.client.get(url).status_code, 401)
        r = self.client.get(url + "?kazanim=1", headers={"Authorization": "Bearer gizli"})
        self.assertEqual(r.status_code, 200)
        veri = r.json()["sonuclar"][0]
        self.assertEqual(veri["numara"], "1")
        self.assertEqual(veri["dogru"], 7)
        self.assertTrue(veri["kazanimlar"])

    def test_satir_duzeltme(self):
        satir = OptikSatir.objects.get()
        r = self.client.post(f"/satir/{satir.pk}/", {"numara": "0042", "kitapcik": "A", "dahil": "on"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.deneme.sonuclar.get().ogrenci.numara, "42")


@OPTIK
class GercekDosyaTest(TestCase):
    """ornek_veri/ klasöründeki gerçek deneme dosyaları varsa uçtan uca doğrular."""

    def test_cinili_saray_deneme1(self):
        dat, xlsx = ORNEK / "deneme1_7sinif.dat", ORNEK / "deneme1_anahtar.xlsx"
        if not (dat.exists() and xlsx.exists()):
            self.skipTest("Gerçek örnek dosyalar yok")
        sablon = OptikSablon.objects.get(ad__startswith="Çinili Saray")
        d = Deneme.objects.create(ad="D1", tarih=datetime.date(2026, 10, 10), sinif_seviyesi="7")
        anahtar_yukle(d, xlsx.read_bytes())
        self.assertEqual(d.toplam_soru, 90)
        y = optik_yukle(d, sablon, "d.dat", dat.read_bytes())
        self.assertEqual(y.satirlar.count(), 20)
        self.assertFalse(y.satirlar.filter(durum=OptikSatir.Durum.HATA).exists())
        self.assertEqual(deneme_hesapla(d), 20)
        # Doğru kitapçık eşleşmesinde her öğrenci 50+ doğru yapar; yanlış eşleşmede ~%25.
        self.assertGreaterEqual(min(d.sonuclar.values_list("dogru", flat=True)), 50)


@OPTIK
class OgrenciAktarimTest(TestCase):
    def test_metin_bicimleri(self):
        from okuma.services.ogrenci_aktar import metni_ayristir

        satirlar, hatalar = metni_ayristir(
            "Numara\tAd Soyad\tSınıf\tŞube\n"
            "0048\tAhmet Selim Gündüz\t7\tA\n"
            "57 Mehmet Kaya 7B\n"
            "60;Ali Veli;7-i\n"
            "abc\n"
        )
        self.assertEqual(
            [(s.numara, s.ad_soyad, s.sinif, s.sube) for s in satirlar],
            [("48", "Ahmet Selim Gündüz", "7", "A"), ("57", "Mehmet Kaya", "7", "B"), ("60", "Ali Veli", "7", "İ")],
        )
        self.assertEqual(len(hatalar), 1)

    def test_excel_ad_soyad_ayri(self):
        from okuma.services.ogrenci_aktar import excel_ayristir

        wb = openpyxl.Workbook()
        wb.active.append(["Öğrenci No", "Adı", "Soyadı", "Sınıf"])
        wb.active.append([48, "Ahmet", "Gündüz", "7/A"])
        buf = io.BytesIO()
        wb.save(buf)
        satirlar, _ = excel_ayristir(buf.getvalue())
        self.assertEqual((satirlar[0].numara, satirlar[0].ad_soyad, satirlar[0].sinif, satirlar[0].sube), ("48", "Ahmet Gündüz", "7", "A"))

    def test_kaydet_gunceller_ve_sonuclara_yansir(self):
        from okuma.models import Ogrenci
        from okuma.services.ogrenci_aktar import OgrenciSatiri, etkilenen_denemeleri_hesapla, ogrencileri_kaydet

        sablon = OptikSablon.objects.get(ad__startswith="Çinili Saray")
        d = Deneme.objects.create(ad="D", tarih=datetime.date(2026, 10, 10), sinif_seviyesi="7")
        anahtar_yukle(d, excel_uret())
        optik_yukle(d, sablon, "x.dat", satir_uret("48", "A", "ABCDBCA", sube=" ").encode("cp1254"))
        deneme_hesapla(d)
        self.assertEqual(d.sonuclar.get().ogrenci.ad_soyad, "")
        sonuc = ogrencileri_kaydet([OgrenciSatiri("48", "Ahmet Gündüz", "7", "B"), OgrenciSatiri("50", "Ali", "7", "A")])
        self.assertEqual((sonuc.eklenen, sonuc.guncellenen), (1, 1))
        etkilenen_denemeleri_hesapla(sonuc.ogrenci_idler)
        s = d.sonuclar.select_related("ogrenci").get()
        self.assertEqual((s.ogrenci.ad_soyad, s.sube), ("Ahmet Gündüz", "B"))
        self.assertEqual(Ogrenci.objects.count(), 2)


@OPTIK
class YonetimSayfaTest(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser("yon", password="x-test-1234")
        self.client.force_login(self.admin)

    def test_sayfalar_aciliyor(self):
        from okuma.models import Ogrenci

        o = Ogrenci.objects.create(numara="1", ad_soyad="A")
        sablon = OptikSablon.objects.first()
        for url in ["/ogrenciler/", "/ogrenciler/toplu/", f"/ogrenciler/{o.pk}/", "/sablonlar/", "/sablonlar/yeni/",
                    f"/sablonlar/{sablon.pk}/", "/kullanicilar/", "/kullanicilar/yeni/", f"/kullanicilar/{self.admin.pk}/",
                    "/deneme/yeni/"]:
            self.assertEqual(self.client.get(url).status_code, 200, url)

    def test_toplu_yapistir(self):
        from okuma.models import Ogrenci

        r = self.client.post("/ogrenciler/toplu/", {"metin": "48\tAhmet\t7\tA\n49\tMehmet\t7\tB"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(Ogrenci.objects.get(numara="49").sube, "B")

    def test_ogretmen_kullanici_yonetemez(self):
        self.client.force_login(User.objects.create_user("ogr", is_staff=True))
        self.assertEqual(self.client.get("/kullanicilar/").status_code, 302)

    def test_sablon_onizleme(self):
        sablon = OptikSablon.objects.first()
        dosya = io.BytesIO(satir_uret("48", "A", "ABCDBCA").encode("cp1254"))
        dosya.name = "o.dat"
        veri = {f: getattr(sablon, f) or "" for f in ["ad", "kodlama", "numara_bas", "numara_uzunluk", "kitapcik_bas",
                "sinif_bas", "sinif_uzunluk", "sube_bas", "sube_uzunluk", "ad_bas", "ad_uzunluk", "cevap_bas", "aciklama"]}
        r = self.client.post(f"/sablonlar/{sablon.pk}/", {**veri, "ornek": dosya, "soru_sayisi": 7, "onizle": "1"})
        self.assertContains(r, "ABCDBCA")


@OPTIK
class DenemeTuruTest(TestCase):
    """Tek optik form, 75 soruluk denemede bloklar kısmen dolu."""

    def setUp(self):
        from okuma.models import DenemeTuru
        from okuma.services.yapi import yapiyi_kur

        self.sablon = OptikSablon.objects.get(ad__startswith="Çinili Saray")
        self.tur75 = DenemeTuru.objects.get(ad="6. Sınıf · 75 soru")
        self.d = Deneme.objects.create(ad="6-75", tarih=datetime.date(2026, 10, 13), sinif_seviyesi="?", tur=self.tur75)
        yapiyi_kur(self.d)

    def test_sekiz_tur_tanimli(self):
        from okuma.models import DenemeTuru

        self.assertEqual(DenemeTuru.objects.count(), 8)
        for t in DenemeTuru.objects.all():
            self.assertEqual(t.optik_uzunluk, 90)
            self.assertIn(t.soru_toplam, (75, 90))
        self.assertIn("İnkılap", DenemeTuru.objects.get(ad="8. Sınıf · 90 soru").dersler.get(sira=2).ders)

    def test_yapi_kurulur(self):
        self.assertEqual(self.d.sinif_seviyesi, "6")
        self.assertEqual([t.soru_sayisi for t in self.d.testler.all()], [15, 10, 10, 10, 15, 15])
        self.assertEqual(self.d.optik_uzunluk, 90)

    def test_75_okuma_blok_ofsetleri(self):
        from okuma.services.anahtar_excel import anahtar_elle

        testler = list(self.d.testler.all())
        anahtarlar = ["A" * 15, "B" * 10, "C" * 10, "D" * 10, "A" * 15, "B" * 15]
        for t, a in zip(testler, anahtarlar):
            anahtar_elle(t, a)
        # Optikte bloklar 20/10/10/10/20/20; 15'lik derslerde son 5 kutu boş.
        cevap = "A" * 15 + " " * 5 + "B" * 10 + "C" * 10 + "D" * 10 + "A" * 15 + " " * 5 + "B" * 15 + " " * 5
        optik_yukle(self.d, self.sablon, "x.dat", satir_uret("12", "A", cevap).encode("cp1254"))
        deneme_hesapla(self.d)
        s = self.d.sonuclar.get()
        self.assertEqual((s.dogru, s.yanlis, s.bos), (75, 0, 0))
        # Tam net LGS'ye oranlanınca puan ~500 olmalı.
        self.assertGreater(s.puan, Decimal("499"))

    def test_excel_tur_yapisina_oturur(self):
        satirlar = []
        for test, ders, n in [("Sözel (TÜR)", "Türkçe", 15), ("Sözel (SOS)", "Sosyal Bilgiler", 10),
                              ("Sözel (DİN)", "Din K.ve A.B.", 10), ("Sözel (İNG)", "İngilizce", 10),
                              ("Sayısal (MAT)", "Matematik", 15), ("Sayısal (FEN)", "Fen", 15)]:
            satirlar += [(test, ders, i, n + 1 - i, "ABCD"[i % 4], "", "k") for i in range(1, n + 1)]
        anahtar_yukle(self.d, excel_uret(satirlar))
        mat = self.d.testler.get(lgs_ders="matematik")
        self.assertEqual((mat.ders, mat.optik_blok, mat.sorular.count()), ("Matematik", 20, 15))

    def test_excel_soru_sayisi_uyusmazsa_hata(self):
        with self.assertRaises(AnahtarHatasi):
            anahtar_yukle(self.d, excel_uret())  # 4+3 soruluk anahtar 6 derslik yapıya uymaz

    def test_elle_anahtar_b_sirasi(self):
        from okuma.services.anahtar_excel import anahtar_elle

        t = self.d.testler.get(lgs_ders="din")
        anahtar_elle(t, "ABCDA BCDAX", "10 9 8 7 6 5 4 3 2 1")
        self.assertEqual(t.sorular.get(no=10).iptal, True)
        self.assertEqual(t.sorular.get(no=1).kitapcik_siralari.get(kitapcik="B").no, 10)
        with self.assertRaises(AnahtarHatasi):
            anahtar_elle(t, "ABC")

    def test_sayfalar_ve_katilmayanlar(self):
        from okuma.models import Ogrenci
        from okuma.services.anahtar_excel import anahtar_elle

        for t in self.d.testler.all():
            anahtar_elle(t, "A" * t.soru_sayisi)
        Ogrenci.objects.create(numara="99", ad_soyad="Gelmeyen Öğrenci", sinif_seviyesi="6", sube="A")
        optik_yukle(self.d, self.sablon, "x.dat", satir_uret("12", "A", "A" * 90).encode("cp1254"))
        deneme_hesapla(self.d)
        self.client.force_login(User.objects.create_superuser("y", password="x-test-1234"))
        pk = self.d.pk
        for url in [f"/deneme/{pk}/", f"/deneme/{pk}/cevap-anahtari/", f"/deneme/{pk}/karneler/", f"/deneme/{pk}/karneler/?sube=",
                    "/sablonlar/", f"/deneme-turleri/{self.tur75.pk}/", f"/deneme/{pk}/duzenle/"]:
            self.assertEqual(self.client.get(url).status_code, 200, url)
        self.assertContains(self.client.get(f"/deneme/{pk}/sonuclar/"), "Gelmeyen Öğrenci")

    def test_yeni_deneme_formu_yapiyi_kurar(self):
        self.client.force_login(User.objects.create_superuser("y2", password="x-test-1234"))
        from okuma.models import DenemeTuru

        tur = DenemeTuru.objects.get(ad="8. Sınıf · 90 soru")
        r = self.client.post("/deneme/yeni/", {"ad": "LGS-1", "tur": tur.pk, "tarih": "2026-10-13", "yanlis_katsayi": 3,
                                                "puan_modu": "lgs_tablo", "lgs_yili": 2024})
        self.assertEqual(r.status_code, 302)
        d = Deneme.objects.get(ad="LGS-1")
        self.assertEqual((d.sinif_seviyesi, d.toplam_soru), ("8", 90))



class AnaSiteBaglantisiTest(TestCase):
    def test_optik_host_kendi_url_yapisina_gider(self):
        r = self.client.get("/", HTTP_HOST="optik.localhost")
        self.assertContains(r, "Giriş yapın")
        self.assertEqual(r["X-Robots-Tag"], "noindex, nofollow, noarchive")
        self.assertEqual(self.client.get("/saglik/", HTTP_HOST="optik.localhost").status_code, 200)

    @override_settings(ROOT_URLCONF="config.optik_urls", OKUMA_ERISIM=None)
    def test_yetkisiz_kullanici_giremez(self):
        self.client.force_login(User.objects.create_user("veli"))  # personel değil
        self.assertEqual(self.client.get("/").status_code, 403)
        self.assertEqual(self.client.get("/ogrenciler/").status_code, 403)

    @override_settings(ROOT_URLCONF="config.optik_urls")
    def test_ana_sitede_kullanici_yonetimi_kapali(self):
        self.client.force_login(User.objects.create_superuser("yon", password="x-test-1234"))
        self.assertEqual(self.client.get("/kullanicilar/").status_code, 404)
        self.assertNotContains(self.client.get("/"), "Kullanıcılar</a>")

    def test_giris_sonrasi_panele_doner(self):
        User.objects.create_superuser("yon2", password="x-test-1234")
        r = self.client.post("/giris/", {"username": "yon2", "password": "x-test-1234"}, HTTP_HOST="optik.localhost")
        self.assertEqual(r.status_code, 302)
        self.assertEqual(r["Location"], "/")


@OPTIK
class RaporTest(TemelTest):
    def setUp(self):
        super().setUp()
        # İki deneme: ikincisinde 2 numara düşüyor.
        self.yukle(satir_uret("1", "A", "ABCDBCA", sube="A"), satir_uret("2", "B", "DCBAACB", sube="B"), satir_uret("3", "A", "AB     ", sube="A"))
        deneme_hesapla(self.deneme)
        self.d2 = Deneme.objects.create(ad="D2", tarih=datetime.date(2026, 11, 1), sinif_seviyesi="7", sablon=self.sablon)
        anahtar_yukle(self.d2, excel_uret())
        optik_yukle(self.d2, self.sablon, "x.dat", "\r\n".join([satir_uret("1", "A", "ABCDBCA"), satir_uret("2", "A", "DDDDDDD")]).encode("cp1254"))
        deneme_hesapla(self.d2)
        self.client.force_login(User.objects.create_superuser("r", password="x-test-1234"))
        self.q = f"?sinif=7&d={self.deneme.pk}&d={self.d2.pk}"

    def test_ortalama_listesi(self):
        from okuma.services.raporlar import ortalama_listesi

        satirlar, ort = ortalama_listesi([self.deneme, self.d2])
        bir = next(s for s in satirlar if s.ogrenci.numara == "1")
        self.assertEqual((bir.sira, bir.deneme_sayisi, bir.dogru), (1, 2, 7))
        self.assertIsNotNone(ort["net"])

    def test_kazanim_yuzde_araligi(self):
        from okuma.services.raporlar import kazanim_yuzde_analizi

        gruplar = kazanim_yuzde_analizi([self.deneme], alt=0, ust=50)
        adlar = {o["ogrenci"].numara for g in gruplar for o in g["ogrenciler"]}
        self.assertIn("3", adlar)
        self.assertNotIn("1", adlar)

    def test_gelisim_dusen_ve_risk(self):
        from okuma.services.raporlar import gelisim_raporu

        g = gelisim_raporu([self.deneme, self.d2])
        iki = next(s for s in g["satirlar"] if s.ogrenci.numara == "2")
        self.assertEqual(iki.kategori, "Düşen")
        self.assertIn("Sert düşüş", iki.riskler)
        uc = next(s for s in g["satirlar"] if s.ogrenci.numara == "3")
        self.assertIn("Devamsızlık", uc.riskler)

    def test_sayfalar_ve_excel(self):
        for ad in ["raporlar/", "raporlar/ortalama/", "raporlar/kazanim/", "raporlar/sinif-konu/", "raporlar/gelisim/"]:
            self.assertEqual(self.client.get("/" + ad).status_code, 200, ad)
            if ad != "raporlar/":
                self.assertEqual(self.client.get(f"/{ad}{self.q}").status_code, 200, ad)
                r = self.client.get(f"/{ad}{self.q}&format=xlsx")
                self.assertEqual(r["Content-Type"][:30], "application/vnd.openxmlformats", ad)
        pk = self.deneme.pk
        test = self.deneme.testler.first()
        self.assertContains(self.client.get(f"/deneme/{pk}/madde/?ders={test.pk}&sube=A"), "Ortalama başarı")
        self.assertEqual(self.client.get(f"/deneme/{pk}/madde/?format=xlsx").status_code, 200)


@OPTIK
@override_settings(OKUMA_ANA_SITE=True)
class AnaSiteAktarimTest(TemelTest):
    """Optik → ana site deneme modülü (taslak), ana sitede yayınla / kaldır."""

    def setUp(self):
        super().setUp()
        from takip.models import EtutHocasi, SinifSube, Talebe

        self.admin = User.objects.create_superuser("aktarim", "a@b.com", "x-test-1234")
        hoca = EtutHocasi.objects.create(ad_soyad="Hoca", user=self.admin)
        sinif = SinifSube.objects.create(sinif="7", sube="A")
        hoca.sorumlu_sinif_subeler.add(sinif)
        self.t1 = Talebe.objects.create(ad_soyad="Ali Veli", sinif_sube=sinif, etut_hocasi=hoca, talebe_no="48")
        self.t2 = Talebe.objects.create(ad_soyad="Can Kaya", sinif_sube=sinif, etut_hocasi=hoca, talebe_no="49")
        self.client.force_login(self.admin)

    def test_esitle_aktar_yayinla(self):
        from takip.deneme_models import DenemeKazanimSonucu, DenemeSinavi, DenemeSoruSonucu
        from okuma.models import Ogrenci

        r = self.client.post("/ogrenciler/ana-siteden/")
        self.assertEqual(r.status_code, 302)
        self.assertEqual(Ogrenci.objects.get(numara="48").dis_kimlik, str(self.t1.pk))

        # 48 ve 49 okunur; 99 ana sitede yok → aktarılmaz.
        self.yukle(satir_uret("48", "A", "ABCDBCA"), satir_uret("49", "B", "DCBA   "), satir_uret("99", "A", "AAAAAAA"))
        deneme_hesapla(self.deneme)
        r = self.client.post(f"/deneme/{self.deneme.pk}/ana-siteye-aktar/")
        self.assertEqual(r.status_code, 302)
        self.deneme.refresh_from_db()
        ana = DenemeSinavi.objects.get(pk=self.deneme.ana_site_deneme_id)
        self.assertEqual(ana.durum, DenemeSinavi.Durum.TASLAK)
        self.assertEqual(ana.sonuclar.count(), 2)
        s1 = ana.sonuclar.get(talebe=self.t1)
        self.assertEqual((s1.toplam_dogru, s1.kurum_sirasi), (7, 1))
        self.assertEqual({b.brans for b in s1.brans_satirlari.all()}, {"turkce", "matematik"})
        self.assertTrue(DenemeKazanimSonucu.objects.filter(deneme=ana, talebe=self.t1).exists())
        # B kitapçığındaki öğrencinin soru numaraları kendi kitapçığına göre
        self.assertEqual(DenemeSoruSonucu.objects.filter(deneme=ana, talebe=self.t2).count(), 7)
        self.assertEqual(
            DenemeSoruSonucu.objects.get(deneme=ana, talebe=self.t2, ders_ad="Türkçe", soru_no=4).sonuc, "dogru"
        )

        # Yeniden aktarım aynı denemeyi günceller
        self.client.post(f"/deneme/{self.deneme.pk}/ana-siteye-aktar/")
        self.assertEqual(DenemeSinavi.objects.filter(pk=ana.pk).count(), 1)
        self.assertEqual(ana.sonuclar.count(), 2)

        # Ana sitede yayınla → aktif; tekrar → taslak
        with override_settings(ROOT_URLCONF="config.urls"):
            self.client.post(f"/yonetim/denemeler/{ana.pk}/yayin/")
            ana.refresh_from_db()
            self.assertEqual(ana.durum, DenemeSinavi.Durum.AKTIF)
            self.client.post(f"/yonetim/denemeler/{ana.pk}/yayin/")
            ana.refresh_from_db()
            self.assertEqual(ana.durum, DenemeSinavi.Durum.TASLAK)



class YayineviAnahtarTest(TestCase):
    """Farklı yayınevi şablonları: başlık aşağıda, farklı sütun adları, 'kod - metin' tek hücrede."""

    def _excel(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append([1, "SINAV ADI"])
        ws.append(["SINAV ADI", None, None, None, None, "LGS BAŞLANGIÇ DENEMESİ"])
        for _ in range(30):
            ws.append([])  # sınav ayarları bloğu
        ws.append(["Test\nNo", None, "Açıklama\nTest Adı vs.", None, "A\nSoru No", None, "B\nSoru No", None, "Doğru\nCevap", None, "Kazanım Kodu veya Kazanım Metni"])
        for test_no, ad, n, kod in [(1, "Sözel (TÜR)", 4, "8.3.2.1"), (5, "Sayısal (MAT)", 3, "8.1.1.1")]:
            for i in range(1, n + 1):
                ws.append([test_no, None, ad, None, i, None, n + 1 - i, None, "ABCD"[i % 4], None, f"{kod}. - Kazanım {i}"])
        ws.append([None, None, "TOPLAM", None, None, None, None, None, None])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_otomatik_algilama(self):
        from okuma.services.anahtar_excel import excel_oku

        sorular = excel_oku(self._excel())
        self.assertEqual(len(sorular), 7)
        ilk = sorular[0]
        self.assertEqual((ilk.test, ilk.a_no, ilk.kitapcik_no, ilk.cevap, ilk.kazanim_kodu, ilk.kazanimlar),
                         ("Sözel (TÜR)", 1, {"B": 4}, "B", "8.3.2.1", ["Kazanım 1"]))

    def test_yukleme_ve_sinif_uyarisi(self):
        d = Deneme.objects.create(ad="X", tarih=datetime.date(2026, 10, 1), sinif_seviyesi="7")
        ozet = anahtar_yukle(d, self._excel())
        self.assertEqual([t.lgs_ders for t in d.testler.all()], ["turkce", "matematik"])
        self.assertIn("8. sınıf", ozet.uyari)

    def test_tablo_yoksa_anlasilir_hata(self):
        wb = openpyxl.Workbook()
        wb.active.append(["Ad", "Soyad"])
        buf = io.BytesIO()
        wb.save(buf)
        with self.assertRaisesMessage(AnahtarHatasi, "Cevap anahtarı tablosu bulunamadı"):
            anahtar_yukle(Deneme.objects.create(ad="Y", tarih=datetime.date(2026, 10, 1), sinif_seviyesi="7"), buf.getvalue())


@OPTIK
class FarkliAnahtarBicimleriTest(TestCase):
    """Kazanımsız, ayrı sayfalı, yan yana ve dizi biçimli anahtarlar; eşleşmesiz B kitapçığı."""

    A = {"Türkçe": "CBDA", "Matematik": "DCD"}
    B = {"Türkçe": "ADBC", "Matematik": "DCD"}  # B, A'nın ters sırası ama eşleştirme verilmemiş

    def setUp(self):
        self.sablon = OptikSablon.objects.get(ad__startswith="Çinili Saray")
        self.d = Deneme.objects.create(ad="5-1", tarih=datetime.date(2026, 10, 13), sinif_seviyesi="5", sablon=self.sablon)

    def _kaydet(self, wb):
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def _ayri_sayfalar(self):
        wb = openpyxl.Workbook()
        for i, (harf, anahtar) in enumerate([("A", self.A), ("B", self.B)]):
            ws = wb.active if i == 0 else wb.create_sheet()
            ws.title = f"Kitapçık - {harf}"
            ws.append([None, "YANIT - 5.SINIF 1"])
            ws.append(["Ders Adı", f"Kitapçık - {harf}", f"Doğru Cevap {harf}"])
            for ders, cevaplar in anahtar.items():
                for no, c in enumerate(cevaplar, 1):
                    ws.append([ders, no, c])
        return self._kaydet(wb)

    def test_ayri_sayfa_eslesmesiz_b(self):
        from okuma.models import KitapcikAnahtari

        ozet = anahtar_yukle(self.d, self._ayri_sayfalar())
        self.assertEqual(ozet.eslesmesiz, ["B"])
        self.assertIn("eşleştirmesi yok", ozet.uyari)
        self.assertEqual(KitapcikAnahtari.objects.get(test__ders="Türkçe", kitapcik="B").cevaplar, "ADBC")
        # B öğrencisi kendi anahtarıyla tam doğru; A öğrencisi de.
        satir_b = satir_uret("2", "B", "ADBCDCD")
        satir_a = satir_uret("1", "A", "CBDADCD")
        optik_yukle(self.d, self.sablon, "x.dat", "\r\n".join([satir_a, satir_b]).encode("cp1254"))
        deneme_hesapla(self.d)
        b = self.d.sonuclar.get(ogrenci__numara="2")
        self.assertEqual((b.dogru, b.a_sirali, b.kitapcik), (7, False, "B"))
        self.assertEqual(self.d.sonuclar.get(ogrenci__numara="1").dogru, 7)
        # Soru analizine yalnız A öğrencisi girer; sayfalar açılır
        maddeler, _ = analiz.madde_analizi(self.d)
        self.assertEqual(maddeler[0].n, 1)
        self.client.force_login(User.objects.create_superuser("k", password="x-test-1234"))
        r = self.client.get(f"/deneme/{self.d.pk}/karne/{b.pk}/")
        self.assertContains(r, "ADBC"[0])
        self.assertContains(self.client.get(f"/deneme/{self.d.pk}/madde/"), "eşleştirmesi olmayan")

    def test_kitapcik_bos_ise_eslesmesiz_anahtarla_tahmin(self):
        anahtar_yukle(self.d, self._ayri_sayfalar())
        from okuma.services.puanlama import harita_kur

        # 10 doğru eşiği küçük testte sağlanmaz; harita B'yi bağımsız mod olarak tanımalı
        self.assertEqual(harita_kur(self.d).kitapcik_modu("B"), "bagimsiz")

    def test_yan_yana_tablolar(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["A KİTAPÇIĞI", None, None, None, "B KİTAPÇIĞI"])
        ws.append(["Ders", "Soru No", "Cevap", None, "Ders", "Soru No", "Cevap"])
        satirlar_a = [(d, i, c) for d, cs in self.A.items() for i, c in enumerate(cs, 1)]
        satirlar_b = [(d, i, c) for d, cs in self.B.items() for i, c in enumerate(cs, 1)]
        for a, b in zip(satirlar_a, satirlar_b):
            ws.append([*a, None, *b])
        from okuma.services.anahtar_excel import anahtar_oku

        okuma_ = anahtar_oku(self._kaydet(wb))
        self.assertEqual(len(okuma_.sorular), 7)
        self.assertEqual(okuma_.bagimsiz["B"]["Türkçe"], "ADBC")

    def test_dizi_bicimi(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["A Kitapçığı"])
        ws.append(["Türkçe", "CBDA"])
        ws.append(["Matematik", "D", "C", "D"])
        from okuma.services.anahtar_excel import anahtar_oku

        okuma_ = anahtar_oku(self._kaydet(wb))
        self.assertEqual([(s.test, s.a_no, s.cevap) for s in okuma_.sorular][:2], [("Türkçe", 1, "C"), ("Türkçe", 2, "B")])
        self.assertEqual(len(okuma_.sorular), 7)

    def test_elle_b_cevaplari(self):
        from okuma.models import KitapcikAnahtari
        from okuma.services.anahtar_excel import anahtar_elle

        anahtar_yukle(self.d, self._ayri_sayfalar())
        t = self.d.testler.get(ders="Matematik")
        anahtar_elle(t, "DCD", "", "CCD")
        self.assertEqual(KitapcikAnahtari.objects.get(test=t, kitapcik="B").cevaplar, "CCD")
        anahtar_elle(t, "DCD", "3 2 1", "")
        self.assertFalse(KitapcikAnahtari.objects.filter(test=t, kitapcik="B").exists())
        self.assertEqual(t.sorular.get(no=1).kitapcik_siralari.get(kitapcik="B").no, 3)


class GercekYayineviDosyalariTest(TestCase):
    """Kullanıcının paylaştığı yayınevi dosyaları (varsa)."""

    DOSYALAR = {
        "5sinif_yanit": "/Users/yahyayazici/.claude/uploads/0a1e3892-85db-4d14-b335-16cc3dfb8d65/ac49d55a-YANIT_-_5.SINIF_1_1.xlsx",
        "8sinif_start": "/Users/yahyayazici/.claude/uploads/0a1e3892-85db-4d14-b335-16cc3dfb8d65/20cefb6e-2026-2027__8__Start_Deneme_kazanim_kodlari.xlsx",
    }

    def test_dosyalar(self):
        from okuma.services.anahtar_excel import anahtar_oku

        for ad, yol in self.DOSYALAR.items():
            if not Path(yol).exists():
                continue
            okuma_ = anahtar_oku(Path(yol).read_bytes())
            if ad == "5sinif_yanit":
                self.assertEqual(len(okuma_.sorular), 75)
                self.assertEqual(sum(len(c) for c in okuma_.bagimsiz["B"].values()), 75)
            else:
                self.assertEqual(len(okuma_.sorular), 90)
                self.assertFalse(okuma_.bagimsiz)
