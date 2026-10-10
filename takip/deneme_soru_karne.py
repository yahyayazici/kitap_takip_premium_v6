"""Karne 2. sayfası — soru bazında doğru / yanlış / boş.

Excel (düz, matris veya uzun), PDF karne ya da bunların zip’i.
Eksi yanlış, artı doğru, B boş. Eşleşen talebenin bu denemedeki
soru satırları yenilenir; diğer talebelerin kayıtları durur.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePosixPath

from django.db import transaction
from openpyxl import load_workbook

from takip.deneme_excel import excel_secim_uyumu, normalize_ad, talebe_eslestir
from takip.deneme_models import DenemeSinavi, DenemeSoruSonucu

DOGRU = DenemeSoruSonucu.Sonuc.DOGRU
YANLIS = DenemeSoruSonucu.Sonuc.YANLIS
BOS = DenemeSoruSonucu.Sonuc.BOS

_DERS = (
    ("t c inkilap tarihi ve ataturkculuk", "T.C. İnkılap Tarihi"),
    ("t c inkilap tarihi", "T.C. İnkılap Tarihi"),
    ("din kulturu ve ahlak bilgisi", "Din Kültürü"),
    ("inkilap tarihi", "İnkılap Tarihi"),
    ("sosyal bilgiler", "Sosyal Bilgiler"),
    ("hayat bilgisi", "Hayat Bilgisi"),
    ("fen bilimleri", "Fen Bilimleri"),
    ("fen bilgisi", "Fen Bilgisi"),
    ("din kulturu", "Din Kültürü"),
    ("ingilizce", "İngilizce"),
    ("matematik", "Matematik"),
    ("turkce", "Türkçe"),
    ("cografya", "Coğrafya"),
    ("geometri", "Geometri"),
    ("biyoloji", "Biyoloji"),
    ("felsefe", "Felsefe"),
    ("inkilap", "İnkılap"),
    ("sosyal", "Sosyal Bilgiler"),
    ("kimya", "Kimya"),
    ("fizik", "Fizik"),
    ("tarih", "Tarih"),
    ("fen", "Fen Bilgisi"),
)

_DERS_TAM = {anahtar: ad for anahtar, ad in _DERS}
_DERS_BITISIK = sorted(
    ((anahtar.replace(" ", ""), ad) for anahtar, ad in _DERS),
    key=lambda cift: len(cift[0]),
    reverse=True,
)

_BASLIK_ROL = {
    "sinif": "sinif",
    "sube": "sinif",
    "sinif sube": "sinif",
    "ad soyad": "ad",
    "adi soyadi": "ad",
    "ogrenci": "ad",
    "isim": "ad",
    "talebe": "ad",
    "ad": "ad",
    "ders": "ders",
    "ders adi": "ders",
    "soru": "soru",
    "soru no": "soru",
    "soru numarasi": "soru",
    "konu": "konu",
    "konular": "konu",
    "kazanim": "konu",
    "sonuc": "sonuc",
    "yn": "sonuc",
    "durum": "sonuc",
}

_KIMLIK_AD = {"ad soyad", "adi soyadi", "ogrenci", "isim", "talebe", "ad", "adi", "soyad", "sinif", "sube"}
_METRIC = {"yn", "yuzde", "net", "sonuc", "durum", "c", "dc", "dogru", "yanlis", "bos", "y", "d", "b"}
_DURAK_TAM = _METRIC | {
    "konular",
    "konu",
    "puan",
    "sira",
    "sirasi",
    "okul",
    "kitapcik",
    "genel",
    "kurum",
    "toplam",
    "ortalama",
    "soru",
    "soru sayisi",
    "analiz",
    "a kitapcik",
}
_DURAK_BAS = (
    "konu analiz",
    "ders analiz",
    "sinava giren",
    "brans basari",
    "genel ort",
    "okul bilgi",
    "sinav tarih",
    "sinav adi",
    "son 2",
)

_DUZ_RE = re.compile(
    r"^(?P<ders>.+?)\s+(?P<no>\d{1,3})(?:\s*[\(\[](?P<konu>[^)\]]+)[)\]])?\s*$"
)
_SINIF_RE = re.compile(r"(\d{1,2})\s*[-–]\s*([A-Za-zÇĞİÖŞÜçğıöşü])")
_HARF_SONUC_RE = re.compile(r"([A-Ea-e])([+\-−–—])")


@dataclass
class SoruImportStats:
    soru_yazilan: int = 0
    eslesen_talebe: int = 0
    eslesmeyen: list[str] = field(default_factory=list)
    ders_sayisi: int = 0
    uyari: list[str] = field(default_factory=list)


@dataclass
class _SoruSatiri:
    sinif: str
    ad: str
    ders: str
    soru_no: int
    konu: str
    sonuc: str


@dataclass
class _Parca:
    text: str
    x0: float
    x1: float
    y: float


def _tr_baslik(metin: str) -> str:
    cevir = str.maketrans({"İ": "i", "I": "ı", "Ş": "ş", "Ğ": "ğ", "Ü": "ü", "Ö": "ö", "Ç": "ç"})
    kucuk = " ".join(metin.translate(cevir).lower().split())

    def kelime(parca: str) -> str:
        if not parca:
            return parca
        ilk = {"i": "İ", "ı": "I", "ş": "Ş", "ğ": "Ğ", "ü": "Ü", "ö": "Ö", "ç": "Ç"}.get(
            parca[0], parca[0].upper()
        )
        return ilk + parca[1:]

    return " ".join(kelime(parca) for parca in kucuk.split())


def _ders_adi(metin: str) -> str | None:
    n = normalize_ad(metin)
    if not n or any(ch.isdigit() for ch in n):
        return None
    if n in _DERS_TAM:
        return _DERS_TAM[n]
    bitisik = n.replace(" ", "")
    for anahtar, ad in _DERS_BITISIK:
        if anahtar == bitisik:
            return ad
    return None


def _ders_basligi(metin: str) -> str | None:
    """8. sınıf karnesi başlığı: «TÜRKÇE 20», «8. SINIF FEN BİLİMLERİ»."""
    temiz = " ".join((metin or "").split())
    if not temiz or _satir_soru(temiz):
        return None
    if len(temiz.split()) > 8:
        return None
    n = normalize_ad(temiz)
    n = re.sub(r"\b\d+\b", " ", n)
    for cop in ("soru", "net", "dersi", "sinif", "sinifi", "lgs", "brans"):
        n = re.sub(rf"\b{cop}\b", " ", n)
    n = re.sub(r"\s+", " ", n).strip()
    if not n:
        return None
    return _ders_adi(n)


_KONU_DERS = (
    ("hava olay", "Fen Bilimleri"),
    ("birey ve toplum", "Sosyal Bilgiler"),
    ("kahraman dog", "İnkılap"),
    ("buyuk taarruz", "İnkılap"),
    ("kuvayi milliye", "İnkılap"),
    ("dunya savas", "İnkılap"),
    ("trablusgarp", "İnkılap"),
    ("milli mucadele", "İnkılap"),
    ("osmanli", "İnkılap"),
    ("buhran", "İnkılap"),
    ("mondros", "İnkılap"),
    ("sakarya", "İnkılap"),
    ("istiklal", "İnkılap"),
    ("kemal", "İnkılap"),
    ("tbmm", "İnkılap"),
    ("sevr", "İnkılap"),
    ("uslu ifade", "Matematik"),
    ("gorsel yorum", "Türkçe"),
    ("fotosentez", "Fen Bilimleri"),
    ("periyodik", "Fen Bilimleri"),
    ("esitsizlik", "Matematik"),
    ("noktalama", "Türkçe"),
    ("fiilimsi", "Türkçe"),
    ("peygamber", "Din Kültürü"),
    ("geometrik", "Matematik"),
    ("olasilik", "Matematik"),
    ("ozdeslik", "Matematik"),
    ("dogrusal", "Matematik"),
    ("cebirsel", "Matematik"),
    ("karekok", "Matematik"),
    ("ataturk", "İnkılap"),
    ("inkilap", "İnkılap"),
    ("paragraf", "Türkçe"),
    ("sozcuk", "Türkçe"),
    ("mevsim", "Fen Bilimleri"),
    ("iklim", "Fen Bilimleri"),
    ("genetik", "Fen Bilimleri"),
    ("basinc", "Fen Bilimleri"),
    ("kader", "Din Kültürü"),
    ("zekat", "Din Kültürü"),
    ("oruc", "Din Kültürü"),
    ("namaz", "Din Kültürü"),
    ("kuran", "Din Kültürü"),
    ("hadis", "Din Kültürü"),
    ("ayet", "Din Kültürü"),
    ("ahlak", "Din Kültürü"),
    ("carpan", "Matematik"),
    ("ucgen", "Matematik"),
    ("yazim", "Türkçe"),
    ("mitoz", "Fen Bilimleri"),
    ("mayoz", "Fen Bilimleri"),
    ("lozan", "İnkılap"),
    ("hello", "İngilizce"),
    ("movie", "İngilizce"),
    ("dna", "Fen Bilimleri"),
)


def konu_ders_adi(metin: str) -> str | None:
    """Ders adı «Sınıf» kalmış satırda konudan dersi bulur."""
    n = normalize_ad(metin)
    if not n:
        return None
    for kalip, ad in _KONU_DERS:
        if " " in kalip:
            kok, son = kalip.rsplit(" ", 1)
            desen = rf"(?:^|\s){re.escape(kok)}\s+{re.escape(son)}\w*(?:\s|$)"
        else:
            desen = rf"(?:^|\s){re.escape(kalip)}"
        if re.search(desen, n):
            return ad
    return None


def _kimlik_yazi(metin: str) -> bool:
    n = normalize_ad(metin)
    if not n:
        return False
    if n in _KIMLIK_AD or n in _BASLIK_ROL:
        return True
    return n.split()[0] in {"sinif", "sinifi", "sube"}


def _durak(metin: str) -> bool:
    n = normalize_ad(metin)
    if not n or n in _DURAK_TAM:
        return True
    if "kitapcik" in n:
        return True
    return any(n.startswith(p) for p in _DURAK_BAS)


def _sonuc_kodu(ham) -> str | None:
    if ham is None:
        return None
    t = str(ham).strip()
    if not t:
        return None
    if t in {"-", "−", "–", "—"}:
        return YANLIS
    if t in {"+", "＋"}:
        return DOGRU
    n = normalize_ad(t)
    if n in {"yanlis", "yanlis yapti", "wrong", "false", "eksi"}:
        return YANLIS
    if n in {"dogru", "true", "arti"}:
        return DOGRU
    if n in {"bos", "b", "blank", "empty"}:
        return BOS
    if n == "y":
        return YANLIS
    if n == "d":
        return DOGRU
    return None


def _soru_no(ham) -> int | None:
    t = str(ham or "").strip()
    if not re.fullmatch(r"\d{1,3}", t):
        return None
    no = int(t)
    if 1 <= no <= 80:
        return no
    return None


def _satir_soru(metin: str) -> tuple[int, str, str] | None:
    temiz = _HARF_SONUC_RE.sub(r"\1 \2", " ".join((metin or "").split()))
    parcalar = temiz.split()
    if not parcalar:
        return None
    no = _soru_no(parcalar[0])
    if no is None:
        return None
    kalan = parcalar[1:]
    sonuc = None
    # Sondaki tek harf cevap anahtarıdır (B = boş değil). Sonuç işareti ayrı durur.
    if kalan and _sonuc_kodu(kalan[-1]) and not re.fullmatch(r"[A-Ea-e]", kalan[-1]):
        sonuc = _sonuc_kodu(kalan[-1])
        kalan = kalan[:-1]
    harfler: list[str] = []
    while kalan and re.fullmatch(r"[A-Ea-e]", kalan[-1]) and len(harfler) < 2:
        harfler.append(kalan.pop().upper())
    if sonuc is None and len(harfler) == 2:
        sonuc = DOGRU if harfler[0] == harfler[1] else YANLIS
    # Boş: yalnız doğru cevap harfi kalır, öğrenci cevabı ve işaret basılmaz.
    if sonuc is None and len(harfler) == 1:
        sonuc = BOS
    # Cevap harfi yoksa satır soru değildir. Sınav adındaki tire yanlış sayılmaz.
    if sonuc is None or not harfler:
        return None
    return no, " ".join(kalan)[:300], sonuc


def _hucre(value) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return " ".join(str(value).split())


def _al(row: list[str], index: int | None) -> str:
    if index is None or index >= len(row):
        return ""
    return row[index]


def _kimlik_hucre(metin: str) -> bool:
    return normalize_ad(metin) in _KIMLIK_AD or normalize_ad(metin) in _BASLIK_ROL


def _duz_hucre(metin: str) -> tuple[str, int, str] | None:
    m = _DUZ_RE.match(metin or "")
    if not m:
        return None
    ders = m.group("ders").strip()
    if not re.search(r"[A-Za-zÇĞİÖŞÜçğıöşüİı]", ders):
        return None
    no = _soru_no(m.group("no"))
    if no is None:
        return None
    return ders, no, (m.group("konu") or "").strip()


def _duz_aday(row: list[str]) -> bool:
    hucreler = [p for p in (_duz_hucre(c) for c in row) if p]
    if not hucreler:
        return False
    bilinen = any(_ders_adi(ders) for ders, _no, _konu in hucreler)
    kimlik = any(_BASLIK_ROL.get(normalize_ad(c)) == "ad" for c in row)
    if bilinen and kimlik:
        return True
    return len(hucreler) >= 2 and (bilinen or kimlik)


def _uzun_roller(row: list[str]) -> dict[str, int] | None:
    roller: dict[str, int] = {}
    for i, cell in enumerate(row):
        rol = _BASLIK_ROL.get(normalize_ad(cell))
        if rol and rol not in roller:
            roller[rol] = i
    if {"ad", "ders", "soru", "sonuc"} <= set(roller):
        return roller
    return None


def _basligi_kes(rows: list[list[str]]) -> list[list[str]]:
    kesilen = 0
    while rows and kesilen < 3:
        dolu = [c for c in rows[0] if c]
        if len(dolu) > 1 or _uzun_roller(rows[0]) or _duz_aday(rows[0]):
            break
        rows = rows[1:]
        kesilen += 1
    return rows


def _kimlik_indeksleri(header_rows: list[list[str]], ilk: int) -> tuple[int | None, int | None]:
    sinif_i = None
    ad_i = None
    for row in header_rows:
        for i in range(min(ilk, len(row))):
            rol = _BASLIK_ROL.get(normalize_ad(row[i]))
            if rol == "sinif":
                sinif_i = i
            elif rol == "ad":
                ad_i = i
    if ad_i is None:
        if ilk >= 2:
            sinif_i = 0 if sinif_i is None else sinif_i
            ad_i = 1
        elif ilk == 1:
            ad_i = 0
        else:
            return None, None
    return sinif_i, ad_i


def _ileri(row: list[str] | None, cols: list[int]) -> list[str]:
    cur = ""
    out: list[str] = []
    for col in cols:
        raw = row[col] if row and col < len(row) else ""
        n = normalize_ad(raw)
        if raw and n not in _METRIC and not _kimlik_hucre(raw):
            cur = raw
        out.append(cur)
    return out


def _uzun_oku(rows: list[list[str]], header_i: int, roller: dict[str, int]) -> list[_SoruSatiri]:
    cikti: list[_SoruSatiri] = []
    for row in rows[header_i + 1 :]:
        ad = _al(row, roller.get("ad"))
        if not ad or normalize_ad(ad) in _KIMLIK_AD:
            continue
        ders = _al(row, roller.get("ders"))
        no = _soru_no(_al(row, roller.get("soru")))
        kod = _sonuc_kodu(_al(row, roller.get("sonuc")))
        if not ders or no is None or not kod:
            continue
        cikti.append(
            _SoruSatiri(
                sinif=_al(row, roller.get("sinif")),
                ad=ad,
                ders=ders,
                soru_no=no,
                konu=_al(row, roller.get("konu"))[:300],
                sonuc=kod,
            )
        )
    return cikti


def _duz_oku(rows: list[list[str]], header_i: int) -> list[_SoruSatiri]:
    header = rows[header_i]
    cols = []
    for i, cell in enumerate(header):
        parsed = _duz_hucre(cell)
        if parsed:
            cols.append((i, parsed))
    if not cols:
        return []
    ilk = cols[0][0]
    sinif_i, ad_i = _kimlik_indeksleri(rows[: header_i + 1], ilk)
    if ad_i is None:
        return []
    cikti: list[_SoruSatiri] = []
    for row in rows[header_i + 1 :]:
        ad = _al(row, ad_i)
        if not ad or normalize_ad(ad) in _KIMLIK_AD:
            continue
        sinif = _al(row, sinif_i)
        for col, (ders, no, konu) in cols:
            kod = _sonuc_kodu(_al(row, col))
            if not kod:
                continue
            cikti.append(_SoruSatiri(sinif, ad, ders, no, konu[:300], kod))
    return cikti


def _veri_baslangici(rows: list[list[str]], soru_idx: int, cols: list[int]) -> tuple[int, list[str]]:
    konu = [""] * len(cols)
    data_start = soru_idx + 1
    for j in range(soru_idx + 1, min(len(rows), soru_idx + 4)):
        row = rows[j]
        qvals = [_al(row, c) for c in cols]
        if any(_sonuc_kodu(v) for v in qvals):
            data_start = j
            break
        norms = [normalize_ad(v) for v in qvals if v]
        if norms and all(n in _METRIC for n in norms):
            data_start = j + 1
            continue
        if any(qvals):
            cur = ""
            for idx, col in enumerate(cols):
                val = _al(row, col)
                if val and normalize_ad(val) not in _METRIC:
                    cur = val
                konu[idx] = cur
            data_start = j + 1
    return data_start, konu


def _ogrenci_satirlari(
    rows: list[list[str]],
    data_start: int,
    cols: list[int],
    dersler: list[str],
    numaralar: list[int],
    konular: list[str],
    sinif_i: int | None,
    ad_i: int,
) -> list[_SoruSatiri]:
    cikti: list[_SoruSatiri] = []
    for row in rows[data_start:]:
        ad = _al(row, ad_i)
        if not ad or normalize_ad(ad) in _KIMLIK_AD:
            continue
        sinif = _al(row, sinif_i)
        for ders, no, konu, col in zip(dersler, numaralar, konular, cols):
            if not ders:
                continue
            kod = _sonuc_kodu(_al(row, col))
            if not kod:
                continue
            cikti.append(_SoruSatiri(sinif, ad, ders, no, (konu or "")[:300], kod))
    return cikti


def _matris_numarali(rows: list[list[str]], soru_idx: int) -> list[_SoruSatiri]:
    soru_row = rows[soru_idx]
    cols = [i for i, c in enumerate(soru_row) if _soru_no(c)]
    if len(cols) < 2:
        return []
    ders_row = rows[soru_idx - 1] if soru_idx else None
    dersler = _ileri(ders_row, cols)
    numaralar = [_soru_no(soru_row[c]) or 0 for c in cols]
    data_start, konular = _veri_baslangici(rows, soru_idx, cols)
    sinif_i, ad_i = _kimlik_indeksleri(rows[:soru_idx], cols[0])
    if ad_i is None:
        return []
    return _ogrenci_satirlari(rows, data_start, cols, dersler, numaralar, konular, sinif_i, ad_i)


def _matris_numarasiz(rows: list[list[str]], ders_idx: int) -> list[_SoruSatiri]:
    ders_row = rows[ders_idx]
    dolu = [
        i
        for i, c in enumerate(ders_row)
        if c and not _kimlik_hucre(c) and normalize_ad(c) not in _METRIC
    ]
    if not dolu:
        return []
    ilk = dolu[0]
    son = ilk
    for row in rows[ders_idx:]:
        for i, c in enumerate(row):
            if c:
                son = max(son, i)
    son = min(son, ilk + 80)
    cols = list(range(ilk, son + 1))
    dersler = _ileri(ders_row, cols)
    numaralar: list[int] = []
    sayac: dict[str, int] = {}
    for ders in dersler:
        anahtar = normalize_ad(ders)
        sayac[anahtar] = sayac.get(anahtar, 0) + 1
        numaralar.append(sayac[anahtar])
    data_start, konular = _veri_baslangici(rows, ders_idx, cols)
    sinif_i, ad_i = _kimlik_indeksleri(rows[: ders_idx + 1], ilk)
    if ad_i is None:
        return []
    return _ogrenci_satirlari(rows, data_start, cols, dersler, numaralar, konular, sinif_i, ad_i)


def _matris_oku(rows: list[list[str]]) -> list[_SoruSatiri]:
    en_iyi = -1
    en_cok = 0
    for i, row in enumerate(rows[:8]):
        adet = sum(1 for c in row if _soru_no(c))
        if adet > en_cok:
            en_cok = adet
            en_iyi = i
    if en_cok >= 2 and en_iyi >= 0:
        return _matris_numarali(rows, en_iyi)
    for i, row in enumerate(rows[:6]):
        if any(_sonuc_kodu(c) for c in row):
            continue
        dolu = [c for c in row if c and not _kimlik_hucre(c)]
        if len(dolu) < 2 or not any(_ders_adi(c) for c in dolu):
            continue
        if any(any(_sonuc_kodu(c) for c in alt) for alt in rows[i + 1 : i + 40]):
            return _matris_numarasiz(rows, i)
    return []


def _tablo_satirlari(rows: list[list[str]]) -> list[_SoruSatiri]:
    rows = _basligi_kes(rows)
    if not rows:
        return []
    for i, row in enumerate(rows[:10]):
        roller = _uzun_roller(row)
        if roller:
            bulunan = _uzun_oku(rows, i, roller)
            if bulunan:
                return bulunan
    for i, row in enumerate(rows[:10]):
        if _duz_aday(row):
            bulunan = _duz_oku(rows, i)
            if bulunan:
                return bulunan
    return _matris_oku(rows)


def _sayfa_satirlari(ws) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in ws.iter_rows(max_row=4000, max_col=120, values_only=True):
        rows.append([_hucre(c) for c in row])
    while rows and not any(rows[-1]):
        rows.pop()
    return rows


def _excel_satirlar(data: bytes) -> list[_SoruSatiri]:
    excel_secim_uyumu()
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"Excel açılamadı: {exc}") from exc
    try:
        for name in wb.sheetnames:
            bulunan = _tablo_satirlari(_sayfa_satirlari(wb[name]))
            if bulunan:
                return bulunan
    finally:
        wb.close()
    return []


def _ad_aday(metin: str) -> bool:
    temiz = " ".join((metin or "").split())
    if len(temiz) < 8 or any(ch.isdigit() for ch in temiz):
        return False
    if _ders_adi(temiz) or _durak(temiz):
        return False
    kelimeler = temiz.split()
    if not 2 <= len(kelimeler) <= 5:
        return False
    n = normalize_ad(temiz)
    if any(k in n.split() for k in ("sinif", "sinav", "okul", "puan", "kitapcik", "analiz")):
        return False
    return all(re.fullmatch(r"[A-Za-zÇĞİÖŞÜçğıöşüİı'.-]+", k) for k in kelimeler)


def _sinif_bul(metin: str) -> str:
    m = _SINIF_RE.search(metin or "")
    if not m:
        return ""
    return f"{int(m.group(1))}-{m.group(2).upper()}"


def _baslik_aday(metin: str) -> bool:
    temiz = " ".join((metin or "").split())
    kelimeler = temiz.split()
    if not 1 <= len(kelimeler) <= 4 or any(ch.isdigit() for ch in temiz):
        return False
    harfler = [ch for ch in temiz if ch.isalpha()]
    if len(harfler) < 3 or not all(ch.isupper() for ch in harfler):
        return False
    if _durak(temiz) or _ad_aday(temiz) or _ders_adi(temiz) or _ders_basligi(temiz):
        return False
    if _kimlik_yazi(temiz):
        return False
    return True


def _kimlik_ust(segmentler: list[_Parca], ad: str, sinif: str) -> tuple[str, str]:
    ders_y = [s.y for s in segmentler if _ders_adi(s.text)]
    tavan = max(ders_y) if ders_y else None

    def ustte(s: _Parca) -> bool:
        return tavan is None or s.y > tavan + 4

    adaylar = [s for s in segmentler if ustte(s) and _ad_aday(s.text)]
    if adaylar:
        ad = " ".join(max(adaylar, key=lambda s: s.y).text.split())
    siniflar = []
    for s in segmentler:
        if not ustte(s):
            continue
        bulunan = _sinif_bul(s.text)
        if bulunan:
            siniflar.append((s.y, bulunan))
    if siniflar:
        sinif = max(siniflar)[1]
    return ad, sinif


def _sutun_sinirlari(ankorlar: list[float], genislik: float) -> list[tuple[float, float]]:
    if not ankorlar:
        return [(0.0, genislik + 5)]
    xs = sorted(set(round(x, 1) for x in ankorlar))
    kumeler = [[xs[0]]]
    for x in xs[1:]:
        if x - kumeler[-1][-1] <= 60:
            kumeler[-1].append(x)
        else:
            kumeler.append([x])
    merkezler = [min(kume) for kume in kumeler]
    sinir = [0.0]
    for a, b in zip(merkezler, merkezler[1:]):
        sinir.append((a + b) / 2)
    sinir.append(genislik + 5)
    return list(zip(sinir, sinir[1:]))


def _ankorlar(segmentler: list[_Parca]) -> list[float]:
    xs = [s.x0 for s in segmentler if _ders_basligi(s.text)]
    if xs:
        return xs
    sirali = sorted(segmentler, key=lambda s: -s.y)
    for i, s in enumerate(sirali):
        if not _baslik_aday(s.text):
            continue
        for q in sirali[i + 1 : i + 6]:
            if q.y < s.y - 2 and abs(q.x0 - s.x0) < 50 and _satir_soru(q.text):
                xs.append(s.x0)
                break
    return xs


def _sutun_sorulari(segmentler: list[_Parca], genislik: float) -> list[_SoruSatiri]:
    cikti: list[_SoruSatiri] = []
    for sol, sag in _sutun_sinirlari(_ankorlar(segmentler), genislik):
        kolon = [s for s in segmentler if sol <= s.x0 < sag]
        kolon.sort(key=lambda s: -s.y)
        ders = ""
        bekleyen = ""
        for s in kolon:
            if _kimlik_yazi(s.text) and not _ders_basligi(s.text):
                bekleyen = ""
                continue
            bilinen = None if s.text[:1].isdigit() else _ders_basligi(s.text)
            if bilinen:
                ders = bilinen
                bekleyen = ""
                continue
            soru = _satir_soru(s.text)
            if soru:
                if bekleyen:
                    yeni = _ders_basligi(bekleyen)
                    if yeni:
                        ders = yeni
                    elif not ders and not _kimlik_yazi(bekleyen):
                        ders = _tr_baslik(bekleyen)
                if ders and not _kimlik_yazi(ders):
                    no, konu, sonuc = soru
                    cikti.append(_SoruSatiri("", "", ders, no, konu, sonuc))
                bekleyen = ""
                continue
            if _baslik_aday(s.text):
                bekleyen = s.text.strip()
                continue
            bekleyen = ""
    return cikti


def _kelimeler(textpage) -> list[_Parca]:
    kelimeler: list[_Parca] = []
    buf: list[tuple[str, float, float, float]] = []

    def flush() -> None:
        if not buf:
            return
        metin = "".join(g[0] for g in buf).strip()
        if metin:
            kelimeler.append(
                _Parca(metin, min(g[1] for g in buf), max(g[2] for g in buf), max(g[3] for g in buf))
            )
        buf.clear()

    once = None
    for i in range(textpage.count_chars()):
        try:
            ch = textpage.get_text_range(i, 1)
            left, _bottom, right, top = textpage.get_charbox(i)
        except Exception:  # noqa: BLE001
            flush()
            once = None
            continue
        if not ch or ch.isspace():
            flush()
            once = None
            continue
        if once is not None:
            gap = left - once[0]
            if abs(top - once[1]) > 6 or gap > 3 or gap < -8:
                flush()
        buf.append((ch, left, right, top))
        once = (right, top)
    flush()
    return kelimeler


def _segmentler(kelimeler: list[_Parca]) -> list[_Parca]:
    satirlar: list[_Parca] = []
    for kelime in kelimeler:
        if satirlar:
            s = satirlar[-1]
            gap = kelime.x0 - s.x1
            if abs(kelime.y - s.y) <= 6 and -2 <= gap <= 18:
                s.text = f"{s.text} {kelime.text}"
                s.x1 = max(s.x1, kelime.x1)
                s.x0 = min(s.x0, kelime.x0)
                s.y = max(s.y, kelime.y)
                continue
        satirlar.append(_Parca(kelime.text, kelime.x0, kelime.x1, kelime.y))
    for s in satirlar:
        s.text = " ".join(s.text.split())
    return satirlar


_CEVAP_RE = re.compile(r"^(?:[A-Ea-e]\s+){1,2}[+\-−–—]$")
_TEK_HARF_RE = re.compile(r"^[A-Ea-e]$")


def _cevap_parcasi(metin: str) -> bool:
    temiz = " ".join((metin or "").split())
    return bool(_CEVAP_RE.match(temiz) or _TEK_HARF_RE.match(temiz))


def _soru_govdesi(metin: str) -> bool:
    parcalar = (metin or "").split()
    return bool(parcalar) and _soru_no(parcalar[0]) is not None


def _soru_satirini_birlestir(segmentler: list[_Parca]) -> list[_Parca]:
    """Konu ile sağdaki cevap sütununu (C C +) aynı satırda birleştirir."""
    if not segmentler:
        return []
    sirali = sorted(segmentler, key=lambda s: (-s.y, s.x0))
    gruplar: list[list[_Parca]] = []
    for parca in sirali:
        if not gruplar or abs(gruplar[-1][-1].y - parca.y) > 3.5:
            gruplar.append([parca])
        else:
            gruplar[-1].append(parca)
    cikti: list[_Parca] = []
    for grup in gruplar:
        grup.sort(key=lambda s: s.x0)
        i = 0
        while i < len(grup):
            bu = grup[i]
            if i + 1 < len(grup) and _cevap_parcasi(grup[i + 1].text):
                diger = grup[i + 1]
                gap = diger.x0 - bu.x1
                tek = bool(_TEK_HARF_RE.match(diger.text.strip()))
                if -2 <= gap <= 160 and (not tek or _soru_govdesi(bu.text)):
                    cikti.append(
                        _Parca(
                            f"{bu.text} {diger.text}",
                            min(bu.x0, diger.x0),
                            max(bu.x1, diger.x1),
                            max(bu.y, diger.y),
                        )
                    )
                    i += 2
                    continue
            cikti.append(bu)
            i += 1
    return cikti


def _sayfa_oku(page, ad: str, sinif: str) -> tuple[list[_SoruSatiri], str, str]:
    textpage = page.get_textpage()
    try:
        segmentler = _soru_satirini_birlestir(_segmentler(_kelimeler(textpage)))
    finally:
        textpage.close()
    ad, sinif = _kimlik_ust(segmentler, ad, sinif)
    satirlar = _sutun_sorulari(segmentler, float(page.get_width() or 595))
    for satir in satirlar:
        satir.ad = ad
        satir.sinif = sinif
    return satirlar, ad, sinif


def _pdf_satirlar(data: bytes) -> list[_SoruSatiri]:
    try:
        import pypdfium2 as pdfium
    except ImportError as exc:
        raise ValueError("PDF okunamadı. Sunucuda pypdfium2 kurulu değil.") from exc
    try:
        doc = pdfium.PdfDocument(data)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"PDF açılamadı: {exc}") from exc
    try:
        ad = ""
        sinif = ""
        hepsi: list[_SoruSatiri] = []
        for i in range(len(doc)):
            page = doc[i]
            try:
                satirlar, ad, sinif = _sayfa_oku(page, ad, sinif)
            finally:
                page.close()
            hepsi.extend(satirlar)
            if len(hepsi) > 20000:
                break
        return hepsi
    finally:
        doc.close()


def _zip_satirlar(data: bytes) -> list[_SoruSatiri]:
    try:
        paket = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ValueError("Zip açılamadı.") from exc
    cikti: list[_SoruSatiri] = []
    with paket:
        uyeler = [m for m in paket.infolist() if not m.is_dir()]
        if len(uyeler) > 400:
            raise ValueError("Zip içinde çok fazla dosya var.")
        for uye in uyeler:
            if "__MACOSX" in uye.filename or uye.file_size > 30_000_000:
                continue
            ad = PurePosixPath(uye.filename).name
            if not ad or ad.startswith("."):
                continue
            if not ad.lower().endswith((".pdf", ".xlsx", ".xlsm")):
                continue
            cikti.extend(_blob_satirlar(ad, paket.read(uye)))
            if len(cikti) > 20000:
                break
    return cikti[:20000]


def _blob_satirlar(ad: str, data: bytes) -> list[_SoruSatiri]:
    low = ad.lower()
    if low.endswith(".zip"):
        return _zip_satirlar(data)
    if low.endswith(".pdf"):
        return _pdf_satirlar(data)
    if low.endswith((".xlsx", ".xlsm")):
        return _excel_satirlar(data)
    raise ValueError(f"Desteklenmeyen dosya: {ad}")


def _dosya_bytes(dosya) -> tuple[str, bytes]:
    ad = str(getattr(dosya, "name", "") or "").replace("\\", "/").split("/")[-1].strip()
    if hasattr(dosya, "seek"):
        try:
            dosya.seek(0)
        except Exception:  # noqa: BLE001
            pass
    data = dosya.read() if hasattr(dosya, "read") else dosya
    if isinstance(data, str):
        data = data.encode()
    return ad, data


def _ders_kaydet(ham: str) -> str:
    return (_ders_basligi(ham) or _ders_adi(ham) or _tr_baslik(ham) or ham).strip()[:120]


def import_soru_karneleri(dosyalar, *, deneme: DenemeSinavi) -> SoruImportStats:
    if dosyalar is None:
        dosyalar = []
    elif not isinstance(dosyalar, (list, tuple)):
        dosyalar = [dosyalar]
    if not dosyalar:
        raise ValueError("Karne dosyası seçin.")

    ham: list[_SoruSatiri] = []
    for dosya in dosyalar:
        ad, data = _dosya_bytes(dosya)
        if not ad:
            raise ValueError("Dosya adı okunamadı.")
        ham.extend(_blob_satirlar(ad, data))
        if len(ham) > 20000:
            ham = ham[:20000]
            break
    if not ham:
        raise ValueError(
            "Karne dosyasında soru sonucu bulunamadı. "
            "PDF karnenin 2. sayfası veya ders ve soru numaralı Excel olmalı."
        )
    if not any(s.ad for s in ham):
        raise ValueError("Karne okundu ama öğrenci adı bulunamadı. Üst bilgide ad soyad olmalı.")

    stats = SoruImportStats()
    onbellek: dict[tuple[str, str], object] = {}
    yazilacak: dict[tuple, DenemeSoruSonucu] = {}
    gorulen: dict[tuple[str, int], int] = {}
    for satir in ham:
        if not satir.ad:
            continue
        anahtar = (satir.ad.strip(), satir.sinif.strip())
        if anahtar not in onbellek:
            onbellek[anahtar] = talebe_eslestir(anahtar[0], anahtar[1])[0]
        talebe = onbellek[anahtar]
        if talebe is None:
            etiket = anahtar[0] + (f" ({anahtar[1]})" if anahtar[1] else "")
            if etiket not in stats.eslesmeyen and len(stats.eslesmeyen) < 40:
                stats.eslesmeyen.append(etiket)
            continue
        ders_ad = _ders_kaydet(satir.ders)
        if not _ders_adi(ders_ad):
            ipucu = konu_ders_adi(satir.konu) or konu_ders_adi(satir.ders)
            if ipucu:
                ders_ad = ipucu
        ders_key = normalize_ad(ders_ad)[:120]
        if not ders_key:
            continue
        sira_anahtar = (ders_key, satir.soru_no)
        if sira_anahtar not in gorulen:
            gorulen[sira_anahtar] = len(gorulen) + 1
        yazilacak[(talebe.id, ders_key, satir.soru_no)] = DenemeSoruSonucu(
            deneme=deneme,
            talebe=talebe,
            ders_ad=ders_ad,
            ders_key=ders_key,
            soru_no=satir.soru_no,
            konu_ad=(satir.konu or "")[:300],
            sonuc=satir.sonuc,
            sira=gorulen[sira_anahtar],
        )

    if not yazilacak:
        raise ValueError(
            "Karneye yazılacak talebe eşleşmedi. İsimler sitedeki talebe adlarıyla aynı olmalı."
        )

    ids = {anahtar[0] for anahtar in yazilacak}
    with transaction.atomic():
        DenemeSoruSonucu.objects.filter(deneme=deneme, talebe_id__in=ids).delete()
        DenemeSoruSonucu.objects.bulk_create(yazilacak.values(), batch_size=500)
    stats.soru_yazilan = len(yazilacak)
    stats.eslesen_talebe = len(ids)
    stats.ders_sayisi = len({kayit.ders_key for kayit in yazilacak.values()})
    if stats.eslesmeyen:
        stats.uyari.append(
            "Eşleşmeyen isimler karneye yazılmadı: " + ", ".join(stats.eslesmeyen[:8])
        )
    return stats
