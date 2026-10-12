"""Kazanım Excel'i → DenemeKazanimSonucu.

1) KonuKazanimDetay (deneme kutuğu):
  satır 1 → ders (merged / forward-fill)
  satır 2 → konu / kazanım
  satır 3 → Yüzde | Net
  satır 4+ → Sınıf | Ad Soyad | değerler

2) KonuKazanimYuzdeAnalizi (alt alta bloklar):
  «Türkçe > T.8.2. KONUŞMA»
  Öğrenci | Şube | Doğru % | Doğru | Yanlış | Boş | Net | Toplam
  satırlar…  (Net / Toplam → net_dogru / net_toplam)

Mevcut Talebe kayıtlarıyla eşleştirir; yeni talebe oluşturmaz.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from django.db import transaction
from openpyxl import load_workbook

from takip.deneme_excel import excel_secim_uyumu, talebe_eslestir
from takip.deneme_models import DenemeKazanimSonucu, DenemeSinavi


def clean_label(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split())


def name_key(value: str) -> str:
    return clean_label(value).casefold()


def parse_decimal(value) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip().replace("%", "").replace(",", ".")
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def parse_net(value) -> tuple[Decimal | None, Decimal | None]:
    if value is None or value == "":
        return None, None
    text = str(value).strip()
    if not text:
        return None, None
    match = re.match(
        r"^\s*(-?\d+(?:[.,]\d+)?)\s*/\s*(-?\d+(?:[.,]\d+)?)\s*$",
        text,
    )
    if not match:
        return None, None
    return parse_decimal(match.group(1)), parse_decimal(match.group(2))


@dataclass
class KazanimImportStats:
    sonuc_yazilan: int = 0
    atlanan_bos: int = 0
    eslesen_talebe: int = 0
    eslesmeyen: list[str] = field(default_factory=list)
    konu_sayisi: int = 0
    uyari: list[str] = field(default_factory=list)


def _forward_fill(values: list) -> list[str | None]:
    filled: list[str | None] = []
    current: str | None = None
    for value in values:
        label = clean_label(value)
        if label:
            current = label
        filled.append(current)
    return filled


def _read_columns(ws) -> list[dict]:
    row1 = [cell.value for cell in ws[1]]
    row2 = [cell.value for cell in ws[2]]
    row3 = [cell.value for cell in ws[3]]
    subjects = _forward_fill(row1)
    topics = _forward_fill(row2)
    columns: list[dict] = []
    for index in range(2, len(row3)):
        metric_raw = clean_label(row3[index]).casefold()
        if metric_raw not in {"yüzde", "yuzde", "net"}:
            continue
        subject = subjects[index] if index < len(subjects) else None
        topic = topics[index] if index < len(topics) else None
        if not subject or not topic:
            continue
        columns.append(
            {
                "col": index + 1,
                "ders": subject,
                "konu": topic,
                "metrik": (
                    "yuzde"
                    if "yuzde" in metric_raw or "yüzde" in metric_raw
                    else "net"
                ),
            }
        )
    return columns


_BLOK_BASLIK = {
    "ogrenci": "ad",
    "öğrenci": "ad",
    "ad soyad": "ad",
    "sube": "sinif",
    "şube": "sinif",
    "sinif": "sinif",
    "sınıf": "sinif",
    "doğru %": "yuzde",
    "dogru %": "yuzde",
    "yüzde": "yuzde",
    "net": "net",
    "toplam": "toplam",
}


def _blok_basligi(row: tuple) -> dict[str, int] | None:
    roller: dict[str, int] = {}
    for i, value in enumerate(row):
        rol = _BLOK_BASLIK.get(name_key(value))
        if rol and rol not in roller:
            roller[rol] = i
    if "ad" in roller and ("yuzde" in roller or "net" in roller):
        return roller
    return None


def _ders_konu(value) -> tuple[str, str] | None:
    label = clean_label(value)
    if ">" not in label:
        return None
    ders, konu = (clean_label(p) for p in label.split(">", 1))
    if not ders or not konu:
        return None
    return ders, konu


def _blok_kayitlari(ws) -> list[dict]:
    """KonuKazanimYuzdeAnalizi: «Ders > Konu» başlığı, altında öğrenci tablosu."""
    kayitlar: list[dict] = []
    ders_konu: tuple[str, str] | None = None
    roller: dict[str, int] | None = None
    for row in ws.iter_rows(values_only=True):
        if not row or not any(v not in (None, "") for v in row):
            continue
        dolu = [v for v in row if v not in (None, "")]
        if len(dolu) == 1:
            bulunan = _ders_konu(dolu[0])
            if bulunan:
                ders_konu = bulunan
                roller = None
            continue
        baslik = _blok_basligi(row)
        if baslik:
            roller = baslik
            continue
        if not ders_konu or not roller:
            continue

        def al(rol: str):
            idx = roller.get(rol)
            return row[idx] if idx is not None and idx < len(row) else None

        ad = clean_label(al("ad"))
        if not ad:
            continue
        kayitlar.append(
            {
                "sinif": clean_label(al("sinif")),
                "ad": ad,
                "ders": ders_konu[0],
                "konu": ders_konu[1],
                "yuzde": parse_decimal(al("yuzde")),
                "net_dogru": parse_decimal(al("net")),
                "net_toplam": parse_decimal(al("toplam")),
            }
        )
    return kayitlar


def _blok_import(ws, *, deneme: DenemeSinavi, stats: KazanimImportStats) -> bool:
    kayitlar = _blok_kayitlari(ws)
    if not kayitlar:
        return False
    stats.konu_sayisi = len({(name_key(k["ders"]), name_key(k["konu"])) for k in kayitlar})
    onbellek: dict[tuple[str, str], object] = {}
    eslesen: set[int] = set()
    yazilacak: dict[tuple, DenemeKazanimSonucu] = {}
    for kayit in kayitlar:
        anahtar = (kayit["ad"], kayit["sinif"])
        if anahtar not in onbellek:
            onbellek[anahtar] = talebe_eslestir(*anahtar)[0]
        talebe = onbellek[anahtar]
        if talebe is None:
            etiket = kayit["ad"] + (f" ({kayit['sinif']})" if kayit["sinif"] else "")
            if etiket not in stats.eslesmeyen and len(stats.eslesmeyen) < 40:
                stats.eslesmeyen.append(etiket)
            continue
        if kayit["yuzde"] is None and kayit["net_dogru"] is None and kayit["net_toplam"] is None:
            stats.atlanan_bos += 1
            continue
        eslesen.add(talebe.id)
        ders_key = name_key(kayit["ders"])
        konu_key = name_key(kayit["konu"])
        yazilacak[(talebe.id, ders_key, konu_key)] = DenemeKazanimSonucu(
            deneme=deneme,
            talebe=talebe,
            ders_key=ders_key[:120],
            konu_key=konu_key[:300],
            ders_ad=kayit["ders"][:120],
            konu_ad=kayit["konu"][:300],
            yuzde=kayit["yuzde"],
            net_dogru=kayit["net_dogru"],
            net_toplam=kayit["net_toplam"],
        )
    if yazilacak:
        DenemeKazanimSonucu.objects.filter(deneme=deneme).delete()
        DenemeKazanimSonucu.objects.bulk_create(yazilacak.values())
        stats.sonuc_yazilan = len(yazilacak)
    stats.eslesen_talebe = len(eslesen)
    return True


def _sutun_import(ws, columns: list[dict], *, deneme: DenemeSinavi, stats: KazanimImportStats) -> None:
    topic_cols: dict[tuple[str, str], dict] = {}
    for col in columns:
        key = (name_key(col["ders"]), name_key(col["konu"]))
        bucket = topic_cols.setdefault(
            key,
            {
                "ders": col["ders"],
                "konu": col["konu"],
                "ders_key": key[0],
                "konu_key": key[1],
                "yuzde_col": None,
                "net_col": None,
            },
        )
        if col["metrik"] == "yuzde":
            bucket["yuzde_col"] = col["col"]
        else:
            bucket["net_col"] = col["col"]
    stats.konu_sayisi = len(topic_cols)

    eslesen: set[int] = set()
    yazilacak: dict[tuple, DenemeKazanimSonucu] = {}
    for row in ws.iter_rows(min_row=4, values_only=False):
        if not row:
            continue
        sinif_ad = clean_label(row[0].value if len(row) > 0 else None)
        ad_soyad = clean_label(row[1].value if len(row) > 1 else None)
        if not ad_soyad:
            continue

        talebe, _tip, _oneriler = talebe_eslestir(ad_soyad, sinif_ad)
        if talebe is None:
            etiket = f"{ad_soyad}" + (f" ({sinif_ad})" if sinif_ad else "")
            if etiket not in stats.eslesmeyen and len(stats.eslesmeyen) < 40:
                stats.eslesmeyen.append(etiket)
            continue
        eslesen.add(talebe.id)

        for meta in topic_cols.values():
            yuzde = None
            net_dogru = None
            net_toplam = None
            if meta["yuzde_col"]:
                idx = meta["yuzde_col"] - 1
                if idx < len(row):
                    yuzde = parse_decimal(row[idx].value)
            if meta["net_col"]:
                idx = meta["net_col"] - 1
                if idx < len(row):
                    net_dogru, net_toplam = parse_net(row[idx].value)
            if yuzde is None and net_dogru is None and net_toplam is None:
                stats.atlanan_bos += 1
                continue
            anahtar = (talebe.id, meta["ders_key"], meta["konu_key"])
            yazilacak[anahtar] = DenemeKazanimSonucu(
                deneme=deneme,
                talebe=talebe,
                ders_key=meta["ders_key"],
                konu_key=meta["konu_key"],
                ders_ad=meta["ders"][:120],
                konu_ad=meta["konu"][:300],
                yuzde=yuzde,
                net_dogru=net_dogru,
                net_toplam=net_toplam,
            )
    if yazilacak:
        DenemeKazanimSonucu.objects.filter(deneme=deneme).delete()
        DenemeKazanimSonucu.objects.bulk_create(yazilacak.values())
        stats.sonuc_yazilan = len(yazilacak)
    stats.eslesen_talebe = len(eslesen)


@transaction.atomic
def import_kazanim_excel(uploaded_file, *, deneme: DenemeSinavi) -> KazanimImportStats:
    stats = KazanimImportStats()
    excel_secim_uyumu()
    wb = load_workbook(uploaded_file, data_only=True, read_only=False)
    try:
        ws = wb[wb.sheetnames[0]]
        columns = _read_columns(ws)
        if columns:
            _sutun_import(ws, columns, deneme=deneme, stats=stats)
        elif not _blok_import(ws, deneme=deneme, stats=stats):
            raise ValueError(
                "Excel'de konu kolonları bulunamadı. KonuKazanimDetay "
                "(ders / konu / Yüzde-Net) ya da KonuKazanimYuzdeAnalizi "
                "(«Ders > Konu» blokları) formatında olmalı."
            )
    finally:
        wb.close()

    if stats.sonuc_yazilan == 0:
        stats.uyari.append("Hiç kazanım satırı yazılmadı; dosya veya eşleşmeleri kontrol edin.")
    if stats.eslesmeyen:
        stats.uyari.append(
            f"{len(stats.eslesmeyen)} isim eşleşmedi (ör. {', '.join(stats.eslesmeyen[:5])})."
        )
    return stats
